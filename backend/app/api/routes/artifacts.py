from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import JSONResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import List, Optional
from datetime import datetime
from app.db.database import get_db
from app.db.models import Artifact, ArtifactStatus, ArtifactType, DocumentClassification, Project
from app.core.config import settings
import os

router = APIRouter(prefix="/artifacts", tags=["artifacts"])


class ArtifactResponse(BaseModel):
    id: int
    artifact_id: str
    project_id: Optional[int]
    name: str
    artifact_type: str
    file_path: Optional[str]
    status: str
    classification: str
    created_by: int
    version: int = 1
    parent_artifact_id: Optional[str] = None
    created_at: datetime
    updated_at: datetime
    preview_url: Optional[str] = None

    class Config:
        from_attributes = True


class ArtifactListResponse(BaseModel):
    artifacts: List[ArtifactResponse]
    total: int


def _preview_url(a) -> Optional[str]:
    try:
        if not os.path.exists(a.file_path):
            return None
        ext = (a.artifact_type.value if hasattr(a.artifact_type, "value") else str(a.artifact_type)).lower()
        if ext in ("csv", "txt", "md", "json", "docx", "xlsx", "pptx"):
            return f"/api/artifacts/{a.artifact_id}/preview"
        return None
    except Exception:
        return None


def _to_response(a) -> ArtifactResponse:
    resp = ArtifactResponse.model_validate(a)
    resp.preview_url = _preview_url(a)
    return resp


@router.get("", response_model=ArtifactListResponse)
async def list_artifacts(
    project_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
):
    query = db.query(Artifact).join(Project, Artifact.project_id == Project.id, isouter=True).filter(
        (Project.company_id == settings.company_id) | (Artifact.company_id == settings.company_id)
    )
    if project_id:
        query = query.filter(Artifact.project_id == project_id)
    if status:
        query = query.filter(Artifact.status == status)
    total = query.count()
    artifacts = query.order_by(Artifact.created_at.desc()).offset(offset).limit(limit).all()
    return ArtifactListResponse(
        artifacts=[_to_response(a) for a in artifacts],
        total=total,
    )


@router.get("/{artifact_id}", response_model=ArtifactResponse)
async def get_artifact(artifact_id: str, db: Session = Depends(get_db)):
    artifact = db.query(Artifact).filter(Artifact.artifact_id == artifact_id).first()
    if not artifact:
        raise HTTPException(status_code=404, detail="Artifact not found")
    return _to_response(artifact)


@router.get("/{artifact_id}/download")
async def download_artifact(artifact_id: str, db: Session = Depends(get_db)):
    artifact = db.query(Artifact).filter(Artifact.artifact_id == artifact_id).first()
    if not artifact:
        raise HTTPException(status_code=404, detail="Artifact not found")
    if not artifact.file_path or not os.path.exists(artifact.file_path):
        raise HTTPException(status_code=404, detail="Artifact file not found")

    with open(artifact.file_path, "rb") as f:
        content = f.read()

    media_type = "application/octet-stream"
    if artifact.artifact_type == ArtifactType.DOCX:
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif artifact.artifact_type == ArtifactType.XLSX:
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif artifact.artifact_type == ArtifactType.PPTX:
        media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif artifact.artifact_type == ArtifactType.PDF:
        media_type = "application/pdf"
    elif artifact.artifact_type == ArtifactType.PY:
        media_type = "text/x-python"
    elif artifact.artifact_type == ArtifactType.JSON:
        media_type = "application/json"
    elif artifact.artifact_type == ArtifactType.CSV:
        media_type = "text/csv"
    elif artifact.artifact_type == ArtifactType.TXT:
        media_type = "text/plain"
    elif artifact.artifact_type == ArtifactType.MD:
        media_type = "text/markdown"

    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{artifact.name}.{artifact.artifact_type.value.lower()}"'}
    )


@router.get("/{artifact_id}/preview")
async def preview_artifact(artifact_id: str, db: Session = Depends(get_db)):
    artifact = db.query(Artifact).filter(Artifact.artifact_id == artifact_id).first()
    if not artifact:
        raise HTTPException(status_code=404, detail="Artifact not found")
    if not artifact.file_path or not os.path.exists(artifact.file_path):
        raise HTTPException(status_code=404, detail="Artifact file not found")

    with open(artifact.file_path, "rb") as f:
        content = f.read()

    ext = (artifact.artifact_type.value if hasattr(artifact.artifact_type, "value") else str(artifact.artifact_type)).lower()

    if ext in ("json",):
        return JSONResponse({"artifact_id": artifact_id, "name": artifact.name, "raw": content.decode("utf-8", errors="replace")})

    if ext in ("csv", "txt", "md"):
        return JSONResponse({
            "artifact_id": artifact_id,
            "name": artifact.name,
            "content": content.decode("utf-8", errors="replace"),
            "mime": "text/plain",
        })

    if ext == "pdf":
        return Response(content=content, media_type="application/pdf")

    if ext in ("docx", "xlsx", "pptx"):
        try:
            summary = _summarize_office(artifact.file_path, ext)
        except Exception as e:
            summary = f"Preview unavailable: {str(e)}"
        return JSONResponse({"artifact_id": artifact_id, "name": artifact.name, "content": summary, "mime": "text/plain"})

    return JSONResponse({"artifact_id": artifact_id, "name": artifact.name, "content": "", "mime": "text/plain"})


def _summarize_office(file_path: str, ext: str) -> str:
    if ext == "docx":
        from docx import Document as DocxDocument
        d = DocxDocument(file_path)
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return "\n".join(lines)
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[Sheet: {ws.title}]")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 30:
                    lines.append("...")
                    break
                lines.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(lines)
    if ext == "pptx":
        from pptx import Presentation
        prs = Presentation(file_path)
        lines = []
        for idx, slide in enumerate(prs.slides, start=1):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text]
            if texts:
                lines.append(f"[Slide {idx}] " + " · ".join(texts))
        return "\n".join(lines)
    return "Preview not available for this format."
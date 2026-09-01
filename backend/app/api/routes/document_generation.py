from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import Dict, Any, Optional
from datetime import datetime
from app.db.database import get_db
from app.db.models import Artifact, ArtifactType, ArtifactStatus, DocumentClassification, Project
from app.core.config import settings
import os
import uuid

router = APIRouter(prefix="/documents/generate", tags=["document-generation"])

ARTIFACT_DIR = "/home/aum/Desktop/ML/Smart India Hackathon/backend/artifacts"
os.makedirs(ARTIFACT_DIR, exist_ok=True)


class GenerateRequest(BaseModel):
    template: str
    data: Dict[str, Any]
    format: str
    project_id: Optional[int] = None
    name: Optional[str] = None


class GenerateResponse(BaseModel):
    artifact_id: str
    file_path: str
    download_url: str


@router.post("", response_model=GenerateResponse)
async def generate_document(request: GenerateRequest, db: Session = Depends(get_db)):
    project = None
    if request.project_id:
        project = db.query(Project).filter(Project.id == request.project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")

    artifact_id = str(uuid.uuid4())
    ext = request.format.lower()
    file_name = f"{artifact_id}.{ext}"
    file_path = os.path.join(ARTIFACT_DIR, file_name)

    try:
        if ext == "docx":
            content = generate_docx(request.template, request.data)
        elif ext == "xlsx":
            content = generate_xlsx(request.template, request.data)
        elif ext == "pptx":
            content = generate_pptx(request.template, request.data)
        elif ext == "pdf":
            content = generate_pdf(request.template, request.data)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported format: {ext}")

        with open(file_path, "wb") as f:
            f.write(content)

        artifact = Artifact(
            artifact_id=artifact_id,
            company_id=settings.company_id,
            project_id=request.project_id,
            name=request.name or f"{request.template}_{datetime.utcnow().strftime('%Y%m%d')}",
            artifact_type=ArtifactType(ext.upper()),
            file_path=file_path,
            status=ArtifactStatus.READY_FOR_REVIEW,
            classification=DocumentClassification.CONFIDENTIAL,
            created_by=1,
            sources=[],
        )
        db.add(artifact)
        db.commit()

        return GenerateResponse(
            artifact_id=artifact_id,
            file_path=file_path,
            download_url=f"/api/artifacts/{artifact_id}/download",
        )

    except Exception as e:
        if os.path.exists(file_path):
            os.remove(file_path)
        raise HTTPException(status_code=500, detail=f"Document generation failed: {str(e)}")


def generate_docx(template: str, data: Dict[str, Any]) -> bytes:
    from docx import Document as DocxDocument
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io

    doc = DocxDocument()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    if template == "approval_note":
        title = data.get("title", "Equipment Integrity Review Approval Note")
        doc.add_heading(title, 0)

        meta = data.get("metadata", {})
        for key, value in meta.items():
            p = doc.add_paragraph()
            p.add_run(f"{key}: ").bold = True
            p.add_run(str(value))

        doc.add_heading("Executive Summary", 1)
        doc.add_paragraph(data.get("summary", "Engineering integrity review is recommended."))

        doc.add_heading("Key Findings", 1)
        findings = data.get("findings", [])
        for finding in findings:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(finding.get("parameter", "")).bold = True
            p.add_run(f": {finding.get('value', '')} (Limit: {finding.get('limit', '')})")

        doc.add_heading("Recommendation", 1)
        doc.add_paragraph(data.get("recommendation", "Immediate engineering review required."))

        doc.add_heading("Sources", 1)
        sources = data.get("sources", [])
        for src in sources:
            doc.add_paragraph(f"• {src.get('title', '')} - {src.get('section', '')}", style='List Bullet')

    elif template == "inspection_report":
        doc.add_heading("Inspection Report", 0)
        for key, value in data.items():
            p = doc.add_paragraph()
            p.add_run(f"{key}: ").bold = True
            p.add_run(str(value))

    else:
        doc.add_heading(template.replace("_", " ").title(), 0)
        for key, value in data.items():
            p = doc.add_paragraph()
            p.add_run(f"{key}: ").bold = True
            p.add_run(str(value))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_xlsx(template: str, data: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    if template == "risk_assessment":
        ws.append(["Parameter", "Value", "Limit", "Status", "Unit"])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        rows = data.get("rows", [])
        for row in rows:
            ws.append([row.get("parameter"), row.get("value"), row.get("limit"), row.get("status"), row.get("unit")])

    elif template == "maintenance_history":
        ws.append(["Date", "Activity", "Technician", "Findings", "Status"])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        rows = data.get("rows", [])
        for row in rows:
            ws.append([row.get("date"), row.get("activity"), row.get("technician"), row.get("findings"), row.get("status")])

    else:
        ws.append(["Field", "Value"])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for key, value in data.items():
            ws.append([key, str(value)])

    for column in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_pptx(template: str, data: Dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    import io

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)

    title_shape = slide.shapes.title
    if title_shape:
        title_shape.text = data.get("title", template.replace("_", " ").title())
        title_shape.text_frame.paragraphs[0].font.size = Pt(28)
        title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x1A, 0x3C, 0x5E)

    if template == "executive_summary":
        content = data.get("content", {})
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(12)
        height = Inches(5.5)

        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True

        for key, value in content.items():
            p = tf.add_paragraph()
            p.text = f"{key}: {value}"
            p.font.size = Pt(14)
            p.space_after = Pt(6)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def generate_pdf(template: str, data: Dict[str, Any]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    import io

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=inch, leftMargin=inch, topMargin=inch, bottomMargin=inch)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle('CustomTitle', parent=styles['Title'], fontSize=18, spaceAfter=12)
    story.append(Paragraph(data.get("title", template.replace("_", " ").title()), title_style))
    story.append(Spacer(1, 12))

    if template == "approval_note":
        story.append(Paragraph("<b>Executive Summary</b>", styles['Heading2']))
        story.append(Paragraph(data.get("summary", "Engineering integrity review is recommended."), styles['Normal']))
        story.append(Spacer(1, 12))

        story.append(Paragraph("<b>Key Findings</b>", styles['Heading2']))
        findings = data.get("findings", [])
        for finding in findings:
            story.append(Paragraph(f"• <b>{finding.get('parameter', '')}:</b> {finding.get('value', '')} (Limit: {finding.get('limit', '')})", styles['Normal']))

        story.append(Spacer(1, 12))
        story.append(Paragraph("<b>Recommendation</b>", styles['Heading2']))
        story.append(Paragraph(data.get("recommendation", "Immediate engineering review required."), styles['Normal']))

    else:
        for key, value in data.items():
            story.append(Paragraph(f"<b>{key}:</b> {value}", styles['Normal']))
            story.append(Spacer(1, 6))

    doc.build(story)
    return buffer.getvalue()
import json
import os
import uuid
from typing import Dict, Any, List, Optional

from app.core.logging import get_logger
from app.core.config import settings

logger = get_logger(__name__)

CHAT_UPLOAD_DIR = "/home/aum/Desktop/ML/Smart India Hackathon/backend/chat_uploads"
os.makedirs(CHAT_UPLOAD_DIR, exist_ok=True)

# Chat-scoped in-memory index keyed by conversation_id -> list of attachment records.
# Chat uploads are temporary and intentionally do NOT pollute the permanent
# knowledge base (RAG). They are only available within the active conversation.
_chat_index: Dict[str, List[Dict[str, Any]]] = {}


def _read_text(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except Exception as e:
        logger.error("chat_attachment_read_failed", path=path, error=str(e))
        return ""


def _analyze_xlsx(path: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[Sheet: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 60:
                lines.append("...")
                break
            lines.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(lines)


def _analyze_csv(path: str) -> str:
    content = _read_text(path)
    return content


def _analyze_docx(path: str) -> str:
    from docx import Document as DocxDocument
    d = DocxDocument(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)


def _analyze_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    lines = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text]
        if texts:
            lines.append(f"[Slide {idx}] " + " · ".join(texts))
    return "\n".join(lines)


def _ocr_image(path: str) -> str:
    try:
        import pytesseract
        from PIL import Image
        text = pytesseract.image_to_string(Image.open(path))
        return text.strip()
    except Exception as e:
        logger.error("ocr_failed", path=path, error=str(e))
        return "(Image uploaded; OCR could not be completed: %s)" % str(e)


def _parse_json(path: str) -> str:
    content = _read_text(path)
    try:
        data = json.loads(content)
        return json.dumps(data, ensure_ascii=False, indent=2)
    except Exception:
        return content


def parse_to_text(path: str, ext: str) -> str:
    ext = ext.lower()
    if ext == "txt":
        return _read_text(path)
    if ext == "csv":
        return _analyze_csv(path)
    if ext == "json":
        return _parse_json(path)
    if ext in ("md", "markdown"):
        return _read_text(path)
    if ext == "xlsx":
        return _analyze_xlsx(path)
    if ext == "docx":
        return _analyze_docx(path)
    if ext == "pptx":
        return _analyze_pptx(path)
    if ext in ("png", "jpg", "jpeg", "tiff", "tif", "bmp"):
        return _ocr_image(path)
    if ext == "pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:
            logger.error("pdf_ocr_fallback_needed", error=str(e))
            return "(PDF uploaded; text not extractable: %s)" % str(e)
    return "(Unsupported file type: .%s)" % ext


def save_chat_attachment(conversation_id: str, filename: str, content: bytes) -> Dict[str, Any]:
    ext = os.path.splitext(filename)[1].lstrip(".").lower() or "txt"
    attachment_id = str(uuid.uuid4())
    safe = f"{attachment_id}_{os.path.basename(filename)}"
    path = os.path.join(CHAT_UPLOAD_DIR, safe)
    with open(path, "wb") as f:
        f.write(content)

    text = parse_to_text(path, ext)

    record = {
        "attachment_id": attachment_id,
        "filename": filename,
        "ext": ext,
        "file_path": path,
        "size": len(content),
        "content": text[:20000],
        "summary": _summarize(text, filename),
    }
    index = _chat_index.setdefault(conversation_id, [])
    index.append(record)
    if len(index) > 20:
        index.pop(0)

    return {
        "attachment_id": attachment_id,
        "filename": filename,
        "ext": ext,
        "size": len(content),
        "summary": record["summary"],
        "content_excerpt": record["content"][:300],
    }


def _summarize(text: str, filename: str) -> str:
    stripped = "\n".join(l for l in text.splitlines() if l.strip())
    first_lines = stripped.splitlines()[:12]
    preview = " · ".join(l.strip()[:80] for l in first_lines if l.strip())
    if not preview:
        preview = "(No readable text extracted)"
    if len(preview) > 500:
        preview = preview[:500] + "..."
    return preview or filename


def get_conversation_attachments(conversation_id: str) -> List[Dict[str, Any]]:
    return _chat_index.get(conversation_id, [])


def clear_conversation_attachments(conversation_id: str) -> None:
    _chat_index.pop(conversation_id, None)


def build_attachment_context(conversation_id: str, db=None) -> str:
    """Build the AI context for uploaded files in a conversation.

    Uses the in-memory index when available; otherwise rebuilds it from the
    conversation's persisted messages so uploaded-file context survives a
    backend restart and reopening the chat from History.
    """
    attachments = get_conversation_attachments(conversation_id)
    if not attachments and db is not None:
        attachments = _restore_attachments_from_messages(conversation_id, db)
    if not attachments:
        return ""
    parts = []
    for i, att in enumerate(attachments, start=1):
        parts.append(
            f"--- Uploaded file {i}: {att['filename']} (type .{att['ext']}) ---\n{att['content'][:6000]}"
        )
    return "\n\n".join(parts)


def _restore_attachments_from_messages(conversation_id: str, db) -> List[Dict[str, Any]]:
    """Re-read persisted attachment metadata from the conversation's user messages."""
    from app.db.repositories import ConversationRepository

    repo = ConversationRepository(db)
    messages = repo.get_messages(conversation_id, limit=100)
    records: List[Dict[str, Any]] = []
    seen: set = set()
    for m in messages:
        for att in (m.attachments or []):
            if not isinstance(att, dict):
                continue
            att_id = att.get("attachment_id")
            if att_id in seen:
                continue
            seen.add(att_id)
            path = att.get("file_path") or ""
            deleted = att.get("deleted", False)
            if deleted or not path or not os.path.exists(path):
                continue
            content = _summarize_restored(path, att.get("ext") or "", att.get("filename") or "file")
            if not content:
                continue
            records.append({
                "attachment_id": att_id,
                "filename": att.get("filename") or "file",
                "ext": att.get("ext") or "txt",
                "file_path": path,
                "size": att.get("size") or os.path.getsize(path),
                "content": content[:20000],
                "summary": att.get("summary") or _summarize(content, att.get("filename") or "file"),
            })
    if records:
        _chat_index[conversation_id] = records
    return records


def _summarize_restored(path: str, ext: str, filename: str) -> str:
    try:
        return parse_to_text(path, ext)
    except Exception:
        return ""

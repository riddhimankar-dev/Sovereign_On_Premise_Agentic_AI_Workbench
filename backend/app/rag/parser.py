import fitz
from typing import List, Dict, Any, Optional
from pathlib import Path
from app.core.logging import get_logger

logger = get_logger(__name__)


def parse_pdf(file_path: Path) -> List[Dict[str, Any]]:
    doc = fitz.open(file_path)
    pages = []
    for page_num, page in enumerate(doc):
        text = page.get_text()
        pages.append({
            "page_number": page_num + 1,
            "text": text,
            "has_text": len(text.strip()) > 0,
        })
    doc.close()
    return pages


def parse_docx(file_path: Path) -> List[Dict[str, Any]]:
    import docx
    doc = docx.Document(file_path)
    full_text = "\n".join([p.text for p in doc.paragraphs])
    return [{
        "page_number": 1,
        "text": full_text,
        "has_text": len(full_text.strip()) > 0,
    }]


def parse_xlsx(file_path: Path) -> List[Dict[str, Any]]:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True)
    pages = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(c) if c is not None else "" for c in row])
        text = "\n".join(["\t".join(r) for r in rows])
        pages.append({
            "page_number": len(pages) + 1,
            "text": text,
            "has_text": len(text.strip()) > 0,
            "sheet_name": sheet_name,
        })
    wb.close()
    return pages


def parse_pptx(file_path: Path) -> List[Dict[str, Any]]:
    from pptx import Presentation
    prs = Presentation(file_path)
    pages = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    texts.append(para.text)
        text = "\n".join(texts)
        pages.append({
            "page_number": i + 1,
            "text": text,
            "has_text": len(text.strip()) > 0,
        })
    return pages


def parse_txt(file_path: Path) -> List[Dict[str, Any]]:
    text = file_path.read_text(encoding="utf-8", errors="ignore")
    return [{
        "page_number": 1,
        "text": text,
        "has_text": len(text.strip()) > 0,
    }]


def parse_document(file_path: Path) -> List[Dict[str, Any]]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(file_path)
    elif suffix == ".docx":
        return parse_docx(file_path)
    elif suffix == ".xlsx":
        return parse_xlsx(file_path)
    elif suffix == ".pptx":
        return parse_pptx(file_path)
    elif suffix in [".txt", ".md", ".json", ".yaml", ".yml", ".csv"]:
        return parse_txt(file_path)
    else:
        logger.warning("unsupported_file_type", path=str(file_path))
        return []
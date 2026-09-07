"""
Spreadsheet and CSV data ingestion for the Sovereign Calculation Engine.

Supports:
- CSV files
- Excel XLSX files
- Column selection
- Numeric validation
- Missing-value detection
- Unit metadata
- Period/date metadata
- Structured datasets for the Calculation Engine
"""

from typing import Any, Dict, List
import csv
import os
import math


class SpreadsheetEngine:

    VERSION = "1.0.0"

    # =========================================================
    # CSV
    # =========================================================

    def read_csv(
        self,
        file_path: str,
        value_column: str = None,
        unit_column: str = None,
        period_column: str = None,
    ) -> Dict[str, Any]:

        if not os.path.exists(file_path):
            raise ValueError(
                f"CSV file not found: {file_path}"
            )

        rows = []

        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:

            reader = csv.DictReader(file)

            if not reader.fieldnames:
                raise ValueError(
                    "CSV file does not contain headers"
                )

            columns = list(reader.fieldnames)

            for row_number, row in enumerate(
                reader,
                start=2
            ):

                clean_row = dict(row)

                # -------------------------------------------------
                # Value column
                # -------------------------------------------------

                if value_column:

                    if value_column not in columns:
                        raise ValueError(
                            f"Column '{value_column}' "
                            f"not found in CSV"
                        )

                    raw_value = row.get(
                        value_column
                    )

                    if raw_value is None or str(
                        raw_value
                    ).strip() == "":

                        raise ValueError(
                            f"Missing value in row "
                            f"{row_number}"
                        )

                    try:
                        value = float(
                            raw_value
                        )
                    except (ValueError, TypeError):

                        raise ValueError(
                            f"Invalid numeric value "
                            f"in row {row_number}: "
                            f"{raw_value}"
                        )

                    if not math.isfinite(value):
                        raise ValueError(
                            f"Non-finite value "
                            f"in row {row_number}"
                        )

                    clean_row["value"] = value

                # -------------------------------------------------
                # Unit
                # -------------------------------------------------

                if unit_column:

                    if unit_column not in columns:
                        raise ValueError(
                            f"Column '{unit_column}' "
                            f"not found in CSV"
                        )

                    clean_row["unit"] = row.get(
                        unit_column
                    )

                # -------------------------------------------------
                # Period
                # -------------------------------------------------

                if period_column:

                    if period_column not in columns:
                        raise ValueError(
                            f"Column '{period_column}' "
                            f"not found in CSV"
                        )

                    clean_row["period"] = row.get(
                        period_column
                    )

                rows.append(clean_row)

        return {
            "source_file": file_path,
            "source_type": "csv",
            "columns": columns,
            "row_count": len(rows),
            "dataset": rows,
            "engine_version": self.VERSION,
        }

    # =========================================================
    # EXCEL
    # =========================================================

    def read_excel(
        self,
        file_path: str,
        sheet_name: str = None,
        value_column: str = None,
        unit_column: str = None,
        period_column: str = None,
    ) -> Dict[str, Any]:

        if not os.path.exists(file_path):
            raise ValueError(
                f"Excel file not found: {file_path}"
            )

        try:
            import openpyxl

        except ImportError:

            raise ValueError(
                "openpyxl is required for Excel files. "
                "Install it using: pip install openpyxl"
            )

        workbook = openpyxl.load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        # ---------------------------------------------------------
        # Select sheet
        # ---------------------------------------------------------

        if sheet_name:

            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet_name}' not found"
                )

            worksheet = workbook[
                sheet_name
            ]

        else:

            worksheet = workbook[
                workbook.sheetnames[0]
            ]

        # ---------------------------------------------------------
        # Read headers
        # ---------------------------------------------------------

        rows_iterator = worksheet.iter_rows(
            values_only=True
        )

        try:
            headers = next(
                rows_iterator
            )

        except StopIteration:

            raise ValueError(
                "Excel sheet is empty"
            )

        headers = [
            str(header).strip()
            if header is not None
            else ""
            for header in headers
        ]

        if not any(headers):
            raise ValueError(
                "Excel sheet does not contain headers"
            )

        # ---------------------------------------------------------
        # Validate requested columns
        # ---------------------------------------------------------

        if value_column and value_column not in headers:

            raise ValueError(
                f"Column '{value_column}' "
                f"not found in Excel sheet"
            )

        if unit_column and unit_column not in headers:

            raise ValueError(
                f"Column '{unit_column}' "
                f"not found in Excel sheet"
            )

        if period_column and period_column not in headers:

            raise ValueError(
                f"Column '{period_column}' "
                f"not found in Excel sheet"
            )

        rows = []

        # ---------------------------------------------------------
        # Read data
        # ---------------------------------------------------------

        for row_number, row_values in enumerate(
            rows_iterator,
            start=2
        ):

            row = {}

            for index, header in enumerate(headers):

                if not header:
                    continue

                value = (
                    row_values[index]
                    if index < len(row_values)
                    else None
                )

                row[header] = value

            # -----------------------------------------------------
            # Value
            # -----------------------------------------------------

            if value_column:

                raw_value = row.get(
                    value_column
                )

                if raw_value is None:

                    raise ValueError(
                        f"Missing value in row "
                        f"{row_number}"
                    )

                if isinstance(
                    raw_value,
                    bool
                ):

                    raise ValueError(
                        f"Boolean value found "
                        f"in row {row_number}"
                    )

                try:

                    value = float(
                        raw_value
                    )

                except (
                    ValueError,
                    TypeError
                ):

                    raise ValueError(
                        f"Invalid numeric value "
                        f"in row {row_number}: "
                        f"{raw_value}"
                    )

                if not math.isfinite(value):

                    raise ValueError(
                        f"Non-finite value "
                        f"in row {row_number}"
                    )

                row["value"] = value

            # -----------------------------------------------------
            # Unit
            # -----------------------------------------------------

            if unit_column:

                row["unit"] = row.get(
                    unit_column
                )

            # -----------------------------------------------------
            # Period
            # -----------------------------------------------------

            if period_column:

                row["period"] = row.get(
                    period_column
                )

            rows.append(row)

        workbook.close()

        return {
            "source_file": file_path,
            "source_type": "xlsx",
            "sheet": worksheet.title,
            "columns": headers,
            "row_count": len(rows),
            "dataset": rows,
            "engine_version": self.VERSION,
        }

    # =========================================================
    # AUTO DETECTION
    # =========================================================

    def read(
        self,
        file_path: str,
        **kwargs
    ) -> Dict[str, Any]:

        extension = (
            os.path.splitext(
                file_path
            )[1]
            .lower()
        )

        if extension == ".csv":

            return self.read_csv(
                file_path,
                **kwargs
            )

        if extension in (
            ".xlsx",
            ".xlsm"
        ):

            return self.read_excel(
                file_path,
                **kwargs
            )

        raise ValueError(
            f"Unsupported spreadsheet format: "
            f"{extension}"
        )


# =============================================================
# SIMPLE HELPER
# =============================================================

def load_spreadsheet(
    file_path: str,
    value_column: str = None,
    unit_column: str = None,
    period_column: str = None,
    sheet_name: str = None,
):

    engine = SpreadsheetEngine()

    extension = os.path.splitext(file_path)[1].lower()

    # CSV does not use sheet_name
    if extension == ".csv":
        return engine.read_csv(
            file_path=file_path,
            value_column=value_column,
            unit_column=unit_column,
            period_column=period_column,
        )

    # Excel can use sheet_name
    if extension in (".xlsx", ".xlsm"):
        return engine.read_excel(
            file_path=file_path,
            value_column=value_column,
            unit_column=unit_column,
            period_column=period_column,
            sheet_name=sheet_name,
        )

    raise ValueError(
        f"Unsupported spreadsheet format: {extension}"
    )
# =============================================================
# CALCULATION DATASET HELPER
# =============================================================

def spreadsheet_to_dataset(
    file_path: str,
    value_column: str = None,
    unit_column: str = None,
    period_column: str = None,
    sheet_name: str = None,
) -> Dict[str, Any]:
    """
    Load CSV/XLSX data and prepare it for the Calculation Engine.

    The returned dataset keeps:
    - period
    - value
    - unit
    - source reference

    No calculation is performed here.
    """

    loaded = load_spreadsheet(
        file_path=file_path,
        value_column=value_column,
        unit_column=unit_column,
        period_column=period_column,
        sheet_name=sheet_name,
    )

    dataset = []

    for row in loaded["dataset"]:
        item = dict(row)

        if "value" not in item:
            raise ValueError(
                "Spreadsheet dataset requires a numeric 'value' field."
            )

        item["source_ref"] = file_path

        dataset.append(item)

    return {
        "source_file": loaded["source_file"],
        "source_type": loaded["source_type"],
        "sheet": loaded.get("sheet"),
        "columns": loaded["columns"],
        "row_count": loaded["row_count"],
        "dataset": dataset,
        "engine_version": loaded["engine_version"],
    }